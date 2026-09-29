"""The watcher end to end: real threads, a fake server, a keyboard the test types into.

Every wait has a deadline, so a deadlock fails the test instead of hanging it.
"""

from __future__ import annotations

import queue
import threading
import time

import openpyxl
import pytest

from resumix_client.stages.approval import Answer, HumanApproval
from resumix_client.stages.calls import Calls
from resumix_client.stages.generation import CvWriter, GenerationWorker
from resumix_client.stages.inbox import InputProcessor
from resumix_client.stages.terminal import Console, Keyboard
from resumix_client.stages.watcher import Watcher
from resumix_client.stages.working import WorkingProcessor
from resumix_client.tracking import build_tracker

from conftest import JD_TEXT, FakeApi, analysis, envelope, make_job

DEADLINE = 10.0


class Typist:
    """A stdin the test types into: ``readline`` blocks, like a terminal's."""

    def __init__(self):
        self._lines: queue.Queue = queue.Queue()

    def readline(self):
        return self._lines.get()

    def type(self, line):
        self._lines.put(line + "\n")

    def close(self):
        self._lines.put("")


class ByFirstLine(FakeApi):
    """Analyses each posting by its first line: ``<company>|<should_apply>``."""

    def analyze(self, text, *, profile, preferences, temperature=None):
        self._record("analyze")
        company, should_apply = text.split("\n", 1)[0].split("|")
        return envelope(analysis(company_name=company, should_apply=should_apply))


class Approves:
    """Says y to every posting it is shown."""

    def __init__(self):
        self.shown: list = []

    def review(self, jd_text, analysis):
        self.shown.append(analysis.company_name)
        return Answer("approve")


def posting(company, should_apply):
    return f"{company}|{should_apply}\n{JD_TEXT}"


def build(api, config, workspace, inbox, keyboard, reviewer):
    console = Console()
    writer = CvWriter(Calls(api), config)
    working = WorkingProcessor(workspace, writer, build_tracker(workspace.root, True), console)
    return Watcher(
        working,
        InputProcessor(inbox, Calls(api), config, working, console,
                       poll_seconds=0.01, settle_seconds=0),
        GenerationWorker(working, writer, console, poll_seconds=0.01),
        HumanApproval(working, reviewer, keyboard, console, poll_seconds=0.01),
        keyboard,
        console,
        join_seconds=0.01,
    )


def in_background(watcher):
    result: dict = {}
    thread = threading.Thread(target=lambda: result.update(code=watcher.run()), daemon=True)
    thread.start()
    return thread, result


def wait_until(condition):
    deadline = time.monotonic() + DEADLINE
    while not condition():
        assert time.monotonic() < deadline, "timed out"
        time.sleep(0.01)


def entries(parent):
    return sorted(p.name for p in parent.glob("*/*"))


@pytest.fixture
def inbox(tmp_path):
    folder = tmp_path / "in"
    folder.mkdir()
    return folder


@pytest.fixture
def typist():
    typist = Typist()
    yield typist
    typist.close()


def test_postings_flow_through_the_four_stages(config, workspace, inbox, typist):
    api = ByFirstLine()
    make_job(workspace.working, "Left_Over", should_apply="YES")  # an earlier run's
    (inbox / "a.txt").write_text(posting("Yes Corp", "YES"))
    (inbox / "b.txt").write_text(posting("Check Corp", "CHECK"))
    (inbox / "c.txt").write_text(posting("No Corp", "NO"))
    (inbox / "d.txt").write_text("too short to be a posting")
    reviewer = Approves()

    thread, result = in_background(build(api, config, workspace, inbox, Keyboard(typist),
                                         reviewer))
    wait_until(lambda: len(entries(workspace.cv)) == 3 and not any(inbox.iterdir()))
    typist.type("q")
    thread.join(DEADLINE)

    assert not thread.is_alive()
    assert result["code"] == 0
    assert reviewer.shown == ["Check Corp"], "only CHECK asks"
    assert entries(workspace.cv) == ["Check_Corp_Head_of_IT", "Left_Over", "Yes_Corp_Head_of_IT"]
    assert entries(workspace.discarded) == ["No_Corp_Head_of_IT"]
    assert [name[18:] for name in entries(workspace.error)] == ["d.txt", "d.txt.log"]
    assert not any(workspace.working.iterdir())
    sheet = openpyxl.load_workbook(workspace.root / "applications.xlsx").active
    assert sheet.cell(row=4, column=1).value is not None, "one row per generated CV"
    assert sheet.cell(row=5, column=1).value is None


def test_q_lets_the_generation_in_hand_finish(config, workspace, inbox, typist):
    class Slow(ByFirstLine):
        def __init__(self):
            super().__init__()
            self.writing = threading.Event()
            self.release = threading.Event()

        def cv_status(self, request_id):
            self.writing.set()
            assert self.release.wait(DEADLINE), "never released"
            return super().cv_status(request_id)

    api = Slow()
    (inbox / "a.txt").write_text(posting("Yes Corp", "YES"))
    thread, result = in_background(build(api, config, workspace, inbox, Keyboard(typist),
                                         Approves()))
    try:
        assert api.writing.wait(DEADLINE)
        typist.type("q")
        time.sleep(0.2)
        assert thread.is_alive(), "q waits for the CV being written"
    finally:
        api.release.set()
    thread.join(DEADLINE)

    assert not thread.is_alive()
    assert result["code"] == 0
    assert entries(workspace.cv) == ["Yes_Corp_Head_of_IT"]


def test_a_failed_job_makes_the_exit_code_1(config, workspace, inbox, typist):
    api = ByFirstLine(fail_on="create_cv")
    (inbox / "a.txt").write_text(posting("Yes Corp", "YES"))
    thread, result = in_background(build(api, config, workspace, inbox, Keyboard(typist),
                                         Approves()))
    wait_until(lambda: len(entries(workspace.error)) == 1)
    typist.type("q")
    thread.join(DEADLINE)
    assert result["code"] == 1


def test_a_crashed_stage_stops_the_watcher_loudly(config, workspace, inbox, typist, capsys):
    watcher = build(ByFirstLine(), config, workspace, inbox, Keyboard(typist), Approves())

    def broken(stop):
        raise RuntimeError("the inbox is gone")

    watcher.inbox.run = broken
    thread, result = in_background(watcher)
    thread.join(DEADLINE)

    assert not thread.is_alive(), "no q needed: a crash stops everything"
    assert result["code"] == 1
    out = capsys.readouterr().out
    assert "the inbox thread stopped" in out and "the inbox is gone" in out
